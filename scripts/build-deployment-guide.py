"""Build the Georgian Word handoff with python-docx (documentation tooling)."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(1.8)
sec.left_margin = sec.right_margin = Cm(2)
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
    style = doc.styles[name]
    style.font.name = 'Segoe UI'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(10.5 if name == 'Normal' else 22 if name == 'Title' else 14)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.12
doc.styles['Normal'].element.get_or_add_rPr().append(OxmlElement('w:lang'))
doc.styles['Normal'].element.rPr[-1].set(qn('w:val'), 'ka-GE')
doc.core_properties.title = 'SmartBraille სერვერზე განთავსების ინსტრუქცია'
doc.core_properties.author = 'SmartBraille'
doc.core_properties.subject = 'biovita.ltd და Windows Server'

def p(text):
    return doc.add_paragraph(text)

def h(text):
    doc.add_heading(text, level=1)

def code(text):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(8)
    par.paragraph_format.line_spacing = 1.02
    for i, line in enumerate(text.splitlines()):
        run = par.add_run(('\n' if i else '') + line)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)

doc.add_paragraph('SmartBraille სერვერზე განთავსების ინსტრუქცია', 'Title')
p('საჯარო მისამართი: https://biovita.ltd')
p('SmartBraille-ისთვის ვიყენებთ biovita.ltd დომენს. ეს დომენი ადრე სხვა პროექტისთვის გვქონდა, ამიტომ ახალი დომენის შეძენის ნაცვლად მას ხელახლა ვიყენებთ. პროექტის სახელია SmartBraille.')
p('ეს ინსტრუქცია განკუთვნილია Windows Server-ის ადმინისტრატორისთვის. მიზანია, საიტი და ხმის დამუშავება მთლიანად სახლის სერვერზე მუშაობდეს, ხოლო ჟიურიმ საიტი ინტერნეტით გახსნას. მისამართი გასაზიარებლად მზად იქნება ქვემოთ აღწერილი დაყენებისა და შემოწმების შემდეგ.')
h('1 სერვერის მომზადება')
p('გამოიყენეთ Windows Server x64. რეკომენდებულია 4 პროცესორის ბირთვი, 16 GB ოპერატიული მეხსიერება და 15 GB თავისუფალი ადგილი SSD-ზე. სერვერი, როუტერი და ინტერნეტი ჩართული უნდა დარჩეს; დეველოპერის კომპიუტერი საჭირო აღარ იქნება.')
p('ყველა მომხმარებლისთვის დააყენეთ Git for Windows, Python 3.12 x64 და Node.js 22.12 ან ახალი ვერსია. საჭიროების შემთხვევაში დააყენეთ Microsoft Visual C++ x64 runtime. პროექტი განათავსეთ C:\\Sites\\smartbraille-ში. ადმინისტრატორის უფლებები საჭიროა ავტომატური გაშვებისა და firewall-ის წესისთვის.')
h('2 დომენის დაკავშირება')
p('biovita.ltd-ის DNS მართვის პანელში დაამატეთ ან შეცვალეთ A ჩანაწერი: Name / Host = @; Value = სახლის როუტერის საჯარო IPv4; TTL = 300 წამი ან პროვაიდერის ნაგულისხმევი მნიშვნელობა. მნიშვნელობაში ჩაწერეთ მხოლოდ IP მისამართი.')
p('ძირეული A ჩანაწერის შეცვლა დომენს ძველი პროექტის ჰოსტინგიდან ამ სერვერზე გადაიყვანს. შეინარჩუნეთ სხვა სერვისების, განსაკუთრებით ელფოსტის MX ჩანაწერები. www.biovita.ltd ცალკე მისამართია და ამ ინსტრუქციით ავტომატურად არ კონფიგურირდება.')
p('როუტერში სერვერს დაუფიქსირეთ შიდა IP და გადაამისამართეთ TCP 80 და TCP 443 სერვერის იმავე პორტებზე. 8000 პორტი დატოვეთ მხოლოდ სერვერის შიდა გამოყენებისთვის. თუ ინტერნეტის პროვაიდერი იყენებს CGNAT-ს, საჭიროა მისგან საჯარო IP ან ცალკე გამართული მუდმივი გვირაბი.')
p('თუ საჯარო IP იცვლება, მოაწყვეთ DNS-ის ავტომატური განახლება. არასწორი AAAA ჩანაწერი გაასწორეთ ან წაშალეთ, თუ IPv6 ამ სერვერამდე არ მიდის. Cloudflare-ის შემთხვევაში პირველ შემოწმებაზე გამოიყენეთ DNS only რეჟიმი.')
code("Resolve-DnsName biovita.ltd -Type A")

doc.add_page_break()
h('3 პროექტის ჩამოტვირთვა და დაყენება')
p('სერვერზე გახსენით PowerShell. Python-ის ქვემოთ მოცემული გზა შეცვალეთ რეალური სისტემური ინსტალაციის გზით; მის სანახავად გამოიყენეთ py -0p.')
code(r'''New-Item -ItemType Directory -Path C:\Sites -Force
Set-Location C:\Sites
git clone https://github.com/darkn3sscontact-stack/smartbraille.git
Set-Location C:\Sites\smartbraille
Set-ExecutionPolicy -Scope Process -ExecutionPolicy RemoteSigned
.\scripts\setup.ps1 -Python 'C:\Program Files\Python312\python.exe' -SkipBrowser
.\scripts\start.ps1''')
p('დაყენება შექმნის Python-ის გარემოს, ჩამოტვირთავს ხმის მოდელებს და ააწყობს საიტს. სხვა კომპიუტერიდან .venv ან node_modules საქაღალდეების გადმოტანა საჭირო არ არის. ჩამოტვირთვის შეფერხებისას კავშირის გასწორების შემდეგ იგივე setup ბრძანება ხელახლა გაუშვით.')
p('სერვერზე გახსენით http://127.0.0.1:8000. სხვა PowerShell ფანჯარაში შეამოწმეთ სერვისი:')
code('Invoke-RestMethod http://127.0.0.1:8000/api/health')
p('დაელოდეთ model_state = ready და voice_available = True მნიშვნელობებს. შემდეგ პირველ ფანჯარაში Ctrl+C-ით გააჩერეთ დროებით გაშვებული სერვერი, რომ ავტომატურ გაშვებას პორტი თავისუფალი დახვდეს.')
p('Natia ხმის მოდელს აქვს პირადი გამოყენების პირობები. ორგანიზაციული განთავსებისთვის საჭიროა შესაბამისი ნებართვა ან სათანადო ლიცენზიის მქონე სხვა ხმა. პირობები მოცემულია docs/THIRD_PARTY_NOTICES.md-ში. -SkipVoice პარამეტრი გამოტოვებს ხმის ჩამოტვირთვას; გახმოვანება შესაბამისი მოდელის დაყენებამდე მიუწვდომელი იქნება.')
h('4 HTTPS და ავტომატური გაშვება')
p('PowerShell გახსენით Run as Administrator რეჟიმში და გაუშვით:')
code(r'''Set-Location C:\Sites\smartbraille
Set-ExecutionPolicy -Scope Process -ExecutionPolicy RemoteSigned
.\scripts\download-caddy.ps1
.\scripts\configure-windows-server.ps1 -Domain 'biovita.ltd'
.\scripts\install-windows-server.ps1 -OpenFirewall''')
p('Caddy უზრუნველყოფს HTTPS-ს და სერტიფიკატის განახლებას, როცა DNS და 80/443 პორტები ხელმისაწვდომია. SmartBraille-App და SmartBraille-HTTPS ამოცანები Windows-ის გადატვირთვის შემდეგ ავტომატურად გაეშვება LOCAL SERVICE ანგარიშით. როუტერის გადამისამართება ცალკე უნდა მოაწყოთ.')
p('თუ IIS ან სხვა პროგრამა უკვე იყენებს 80/443 პორტებს, ჯერ მოაგვარეთ პორტების კონფლიქტი. .env, სერტიფიკატები და deploy/runtime საქაღალდის შიგთავსი საჯარო რეპოზიტორიაში არ ატვირთოთ.')

doc.add_page_break()
h('5 საჯარო მისამართის შემოწმება')
code(r'''Get-ScheduledTask -TaskName 'SmartBraille-*'
Invoke-RestMethod http://127.0.0.1:8000/api/health
.\.venv\Scripts\python.exe scripts/verify-public-url.py https://biovita.ltd''')
p('ტელეფონზე გამორთეთ Wi-Fi და მობილური ინტერნეტით გახსენით https://biovita.ltd. შეამოწმეთ მთავარი გვერდი, ლაბორატორია, ანბანი, სავარჯიშო და მოწყობილობა; პირდაპირ გახსენით და განაახლეთ /device გვერდიც.')
p('გამოსცადეთ სიტყვების დაკვრა, სოლენოიდის ხმა, მიკროფონის ნებართვა და ქართული სიტყვის ამოცნობა. შეამოწმეთ შედეგების ექსპორტი და მოწყობილობის ნაწილებზე დაჭერა. ხმას პირველი მომხმარებლის მოქმედება სჭირდება; მიკროფონს სჭირდება სანდო HTTPS და ბრაუზერის ნებართვა.')
p('გადატვირთეთ სერვერი და დარწმუნდით, რომ ორივე ამოცანა იწყება და მოდელები მზადდება. ავტომატური შემოწმება რეალურ ტელეფონზე მიკროფონისა და ხმის მოსმენას ვერ ჩაანაცვლებს.')
p('მხოლოდ წარმატებული გარე შემოწმების შემდეგ .env ფაილში ჩაწერეთ PUBLIC_URL_VERIFIED=true და გადატვირთეთ აპის ამოცანა. ეს გაააქტიურებს საჯარო ბმულისა და QR-ის გამოყენებას.')
code('Stop-ScheduledTask -TaskName SmartBraille-App\n# Wait until the task is no longer Running.\nStart-ScheduledTask -TaskName SmartBraille-App')
p('ჟიურისთვის გასაგზავნი ბმულია https://biovita.ltd. ბმულს არ დაუმატოთ :8000. ვიზიტორებს GitHub-ის ანგარიში და პროგრამების დაყენება არ სჭირდებათ. GitHub-ზე ინახება კოდი; საიტი მუშაობს თქვენს Windows Server-ზე.')
h('6 ხარვეზები და მუშაობის დასრულება')
p('თუ საიტი არ იხსნება, შეამოწმეთ DNS, საჯარო IP, როუტერის 80/443 წესები და firewall. 502 შეცდომისას შეამოწმეთ SmartBraille-App და ლოკალური /api/health. ლოგებია deploy/runtime/logs/app.log და caddy.log. 403 პასუხისას შეამოწმეთ .env-ის ALLOWED_ORIGINS-ში https://biovita.ltd.')
p('ჰოსტინგი 30 დღის შემდეგ თავისით არ ითიშება. დემონსტრაციის დასრულებისას ადმინისტრატორის PowerShell-ში გაუშვით:')
code(r''' 'SmartBraille-HTTPS','SmartBraille-App' | ForEach-Object {
    Disable-ScheduledTask -TaskName $_
    Stop-ScheduledTask -TaskName $_
}
Remove-NetFirewallRule -Name SmartBraille-Web'''.lstrip())
p('როუტერშიც მოხსენით ამ პროექტისთვის დამატებული გადამისამართება. კოდი და მოდელები დარჩება სერვერზე. სრული ინსტრუქცია, განახლების ნაბიჯები და ლიცენზიები ხელმისაწვდომია რეპოზიტორიაში:')
p('https://github.com/darkn3sscontact-stack/smartbraille#readme')

out = ROOT / 'output/docx/SmartBraille_Biovita_Windows_Server_KA.docx'
out.parent.mkdir(parents=True, exist_ok=True)
doc.save(out)
print(out)
